"""規定フォルダの Word / Excel を読み、規定ナビ用のデータを作る。

使い方:
  python build_data.py <規定フォルダ> <出力フォルダ>

出力:
  規定データ.xlsx … Googleスプレッドシートに取り込む用(シート: 規定 / 条文 / 表 / 言い換え / よくある質問)
  data.json        … 手元で画面を確認する用(GitHubには上げない)

規定の公開範囲(一般/役員)は EXEC_ONLY で決める。取り込んだ後にスプレッドシートの
「規定」シートの「公開範囲」列を直しても反映される(ただし次に xlsx を取り込み直すと戻る)。
"""
import sys
import re
import json
import datetime
import pathlib

import docx
import openpyxl
from openpyxl.styles import Font, Alignment

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from docx_numbered import iter_blocks  # noqa: E402

FW2HW = str.maketrans("０１２３４５６７８９", "0123456789")

# 規程一覧表の大区分(管理No.の範囲)
CATEGORIES = [
    ("A", "基本規程", 1, 8),
    ("B", "組織関連規程", 9, 15),
    ("C", "コンプライアンス関連規程", 16, 23),
    ("D", "総務関連規程", 24, 34),
    ("E", "業務関連規程", 35, 40),
    ("F", "経理関連規程", 41, 50),
    ("G", "人事・労務関連規程", 51, 57),
]

# 役員向けの合言葉でだけ見られるもの(それ以外は一般でも見られる)
EXEC_ONLY = {
    "01", "02", "03", "04", "05", "06", "07", "08", "08-1",
    "12", "15", "21", "21-1", "21-2", "23",
    "35", "36", "41", "42", "43", "44", "45", "48",
}

MAX_SECTION = 15000      # 1つの条文セルの最大文字数(スプレッドシートの上限は5万字)
TABLE_CHUNK = 40         # Excelの表を何行ずつ1つの検索単位にするか

BOILERPLATE = [
    (r"株式会社ヘルシーホーム（以下、?「(当社|会社)」という。?）", "当社"),
    (r"及び(その)?子会社もしくは関連会社(に適用され、)?", ""),
]

# 目的の条からうまく要約が作れないものは手で書く
SUMMARY_OVERRIDE = {
    "01": "会社の商号・事業目的・株式・株主総会・取締役などを定めた、会社の根本ルール。",
    "23": "取締役・社員の職務が法令や定款に沿って行われるための、内部統制の基本方針。",
    "51": "正社員の採用・勤務時間・休日・休暇・服務規律・懲戒など、働くうえでの基本ルール。",
    "57": "退職金の支給対象・計算方法・支払いについて定めたもの。",
}


def norm_space(s):
    return re.sub(r"[\s　]+", "", s)


def file_no(name):
    m = re.match(r"^(\d+(?:-\d+)?)_", name)
    return m.group(1) if m else ""


def category_of(no):
    n = int(no.split("-")[0])
    for code, label, lo, hi in CATEGORIES:
        if lo <= n <= hi:
            return f"{code} {label}"
    return ""


def read_master(folder):
    """08-1_規程一覧表 から 管理No. → (規程名, 所管部署) を読む"""
    out = {}
    f = next(folder.glob("08-1_*.xlsx"), None)
    if not f:
        return out
    ws = openpyxl.load_workbook(str(f), data_only=True, read_only=True).worksheets[0]
    for r in ws.iter_rows(values_only=True):
        cells = [("" if v is None else str(v).strip()) for v in r]
        vals = [c for c in cells if c]
        if len(vals) >= 3 and re.match(r"^\d+(-\d+)?$", vals[0]):
            no = vals[0] if "-" in vals[0] else vals[0].zfill(2)
            out[no] = (vals[1], vals[2])
    return out


def clean_name(stem):
    s = re.sub(r"^\d+(-\d+)?_", "", stem)
    s = re.sub(r"_?20\d{6}$", "", s)
    s = re.sub(r"（雛）", "", s)
    s = re.sub(r"（ヘルシーホーム）", "", s)
    return s.strip("_ ")


# ---------------------------------------------------------------- Word

def is_chapter(label, text):
    if re.match(r"^第.+[章節]$", label):
        return True
    if not label and re.match(r"^第[０-９0-9一二三四五六七八九十]+[章節][\s　]", text + " "):
        return True
    return not label and norm_space(text) in ("総則", "附則", "補則", "雑則", "付則")


def is_heading(label, text):
    return not label and re.match(r"^[（(][^（）()]{1,40}[）)]$", text) is not None


def parse_docx(path):
    blocks = list(iter_blocks(docx.Document(str(path))))
    has_auto = any(b[0] == "p" and re.match(r"^第.+条$", b[1]) for b in blocks)
    has_literal = any(b[0] == "p" and not b[1] and re.match(r"^第[０-９0-9]+条[\s　]", b[2])
                      and not b[3].lower().startswith("toc") for b in blocks)
    no_article = not has_auto and not has_literal

    sections = []
    chapter = ""
    chap = ""
    cur = None
    started = False

    def new_section(title="", number=""):
        nonlocal cur
        cur = {"chapter": chapter, "article": number, "title": title, "lines": []}
        sections.append(cur)
        return cur

    for i, b in enumerate(blocks):
        if b[0] == "table":
            rows = [[c for c in r] for r in b[1] if any(c for c in r)]
            if not rows:
                continue
            # 先頭の「株式会社ヘルシーホーム ○○規程」の飾り枠は飛ばす
            if not started and len(rows) <= 2:
                continue
            if cur is None:
                new_section("表")
            for r in rows:
                cur["lines"].append("｜" + "｜".join(c.replace("\n", " ") for c in r) + "｜")
            continue

        _, label, text, style = b
        if style.lower().startswith("toc") or norm_space(text) == "目次":
            continue

        if is_chapter(label, text):
            name = (label + "　" + text).strip("　") if label else text
            name = re.sub(r"[\s　]+", "　", name)
            if re.match(r"^第.+節", name) and chap:
                chapter = chap + "　＞　" + name     # 節は章の下に付ける
            else:
                chap = chapter = name
            started = True
            cur = None
            continue

        if is_heading(label, text):
            new_section(title=text.strip("（）()"))
            started = True
            continue

        m_auto = re.match(r"^第.+条$", label)
        m_lit = None if has_auto else re.match(r"^(第[０-９0-9]+条)[\s　]+(.*)$", text) if not label else None
        if m_auto or m_lit:
            number = label if m_auto else m_lit.group(1)
            body = text if m_auto else m_lit.group(2)
            # 「第1条（適用範囲）」のように番号の段落に見出しだけが入っている形
            if m_auto and re.match(r"^[（(][^（）()]{1,40}[）)]$", body):
                new_section(title=body.strip("（）()"), number=number)
                started = True
                continue
            if cur is not None and not cur["article"] and not cur["lines"]:
                cur["article"] = number
                cur["chapter"] = chapter
            else:
                new_section(number=number)
            cur["lines"].append(body)
            started = True
            continue

        if no_article and re.match(r"^[０-９0-9]+[．.]$", label):
            new_section(title=text)
            started = True
            continue

        if re.match(r"^別[紙表]", text) and len(text) <= 40:
            new_section(title=text)
            started = True
            continue

        if not started:
            # 表題・会社名などの前置き。長い文章(前文)だけ残す
            if len(text) < 40:
                continue
            new_section(title="前文")
            started = True

        if cur is None:
            new_section()
        if label:
            body = re.sub(r"^[.．、]\s*", "", text)
            cur["lines"].append(f"{label}　{body}" if body else label)
        else:
            cur["lines"].append(text)

    out = []
    for s in sections:
        body = "\n".join(s["lines"]).strip()
        if not body and not s["title"]:
            continue
        # 長すぎる条は分ける
        parts = split_text(body, MAX_SECTION)
        for k, p in enumerate(parts):
            t = s["title"] + (f"（{k + 1}/{len(parts)}）" if len(parts) > 1 else "")
            out.append({"chapter": s["chapter"], "article": s["article"], "title": t, "body": p})
    return out


def split_text(body, limit):
    if len(body) <= limit:
        return [body]
    parts, buf = [], []
    size = 0
    for line in body.split("\n"):
        if size + len(line) > limit and buf:
            parts.append("\n".join(buf))
            buf, size = [], 0
        buf.append(line)
        size += len(line) + 1
    if buf:
        parts.append("\n".join(buf))
    return parts


def summary_of(sections):
    for s in sections:
        if s["title"] in ("目的", "目　的"):
            text = s["body"].split("\n")[0]
            for pat, rep in BOILERPLATE:
                text = re.sub(pat, rep, text)
            text = re.sub(r"^(この|本)(規程|規則|マニュアル)は、?", "", text)
            m = re.search(r"^(.+?目的と(する|します)。)", text)
            text = m.group(1) if m else text
            return text[:160]
    return ""


# ---------------------------------------------------------------- Excel

def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y/%m/%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(path):
    wb = openpyxl.load_workbook(str(path), data_only=True)
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or "レビュー" in ws.title:
            continue
        rows = [[cell_str(c.value) for c in r] for r in ws.iter_rows()]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        keep = [j for j in range(width) if any(r[j] for r in rows)]
        rows = [[r[j] for j in keep] for r in rows]
        sheets.append({"sheet": ws.title, "rows": rows})
    return sheets


# ---------------------------------------------------------------- 検索の補助(初期値。スプレッドシートで自由に直せる)

SYNONYMS = [
    ("有給", "年次有給休暇,有休,年休"),
    ("有休", "年次有給休暇,有給,年休"),
    ("産休", "産前産後休業,産前,産後"),
    ("育休", "育児休業"),
    ("介護休み", "介護休業,介護休暇"),
    ("残業", "時間外労働,時間外勤務,休日労働"),
    ("給料", "賃金,給与,基本給"),
    ("給与", "賃金,給料,基本給"),
    ("ボーナス", "賞与"),
    ("退職金", "退職金"),
    ("やめる", "退職,辞職"),
    ("辞める", "退職,辞職"),
    ("クビ", "解雇"),
    ("交通費", "通勤手当,通勤"),
    ("通勤費", "通勤手当"),
    ("出張", "出張,旅費"),
    ("経費", "経費,立替,精算"),
    ("パワハラ", "パワーハラスメント,ハラスメント"),
    ("セクハラ", "セクシャルハラスメント,ハラスメント"),
    ("マタハラ", "マタニティハラスメント,ハラスメント"),
    ("通報", "内部通報,公益通報,相談窓口"),
    ("結婚", "結婚,婚姻,慶弔"),
    ("出産", "出産,慶弔,産前産後"),
    ("葬式", "忌引,弔慰,慶弔,死亡"),
    ("忌引き", "忌引,慶弔"),
    ("病気", "傷病,私傷病,病気休職"),
    ("休職", "休職"),
    ("懲戒", "懲戒,制裁,けん責,減給"),
    ("社用車", "社有車,車両"),
    ("社有車", "車両,社用車"),
    ("事故", "事故,損傷,賠償"),
    ("ハンコ", "印章,押印,印鑑"),
    ("印鑑", "印章,押印"),
    ("個人情報", "個人情報,個人データ"),
    ("マイナンバー", "個人番号,特定個人情報"),
    ("パソコン", "情報システム,PC,端末"),
    ("パスワード", "パスワード,ID"),
    ("決裁", "決裁,稟議,承認,職務権限"),
    ("稟議", "稟議,決裁"),
    ("クレーム", "クレーム,苦情"),
    ("反社", "反社会的勢力"),
    ("インサイダー", "インサイダー取引,内部者取引,重要事実"),
    ("株", "株式,インサイダー取引"),
    ("健康診断", "健康診断,健康情報"),
    ("テレワーク", "在宅勤務,テレワーク"),
    ("副業", "副業,兼業"),
    ("遅刻", "遅刻,早退,欠勤"),
    ("休み", "休日,休暇"),
]

FAQ = [
    ("有給休暇を取りたい", "年次有給休暇", "一般"),
    ("残業・休日出勤", "時間外労働", "一般"),
    ("遅刻・早退・欠勤", "遅刻", "一般"),
    ("給与・手当", "手当", "一般"),
    ("賞与", "賞与", "一般"),
    ("交通費", "通勤手当", "一般"),
    ("結婚・出産・お葬式", "慶弔", "一般"),
    ("育児・介護の休み", "育児休業", "一般"),
    ("病気・ケガで休む", "休職", "一般"),
    ("退職したい", "退職", "一般"),
    ("退職金", "退職金", "一般"),
    ("ハラスメントの相談", "ハラスメント 相談", "一般"),
    ("内部通報", "内部通報", "一般"),
    ("社有車の運転・事故", "車両 事故", "一般"),
    ("個人情報の扱い", "個人情報", "一般"),
    ("パソコン・システム", "情報システム", "一般"),
    ("ハンコ・押印", "印章", "一般"),
    ("決裁・稟議", "稟議", "一般"),
    ("クレーム対応", "クレーム", "一般"),
    ("取締役会", "取締役会", "役員"),
    ("予算", "予算", "役員"),
]


# ---------------------------------------------------------------- 出力

def main():
    src = pathlib.Path(sys.argv[1])
    out_dir = pathlib.Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)
    master = read_master(src)

    regs, secs, tables = [], [], []
    files = sorted(list(src.glob("*.docx")) + list(src.glob("*.xlsx")),
                   key=lambda p: [int(x) for x in file_no(p.name).split("-")])
    for f in files:
        if f.name.startswith("~$"):
            continue
        no = file_no(f.name)
        key = no if "-" in no else no.zfill(2)
        name, dept = master.get(key, (clean_name(f.stem), ""))
        name = re.sub(r"（別[紙表]：.+?）", "", name)
        rid = key
        reg = {
            "id": rid, "no": key, "name": name, "category": category_of(key), "dept": dept,
            "scope": "役員" if key in EXEC_ONLY else "一般",
            "summary": "", "file": f.name, "kind": "表" if f.suffix == ".xlsx" else "規程",
        }
        if f.suffix == ".docx":
            ss = parse_docx(f)
            reg["summary"] = SUMMARY_OVERRIDE.get(key) or summary_of(ss)
            for k, s in enumerate(ss, 1):
                secs.append({"reg": rid, "seq": k, **s})
        else:
            for sh in parse_xlsx(f):
                for k, r in enumerate(sh["rows"], 1):
                    tables.append({"reg": rid, "sheet": sh["sheet"], "row": k, "cells": r})
            reg["summary"] = {"08-1": "社内規程の一覧(所管部署・決裁機関・制定日など)",
                              "09-1": "会社の組織図",
                              "11-1": "決裁項目ごとに、誰が起案・承認・決裁するかの一覧",
                              "24-1": "書類ごとの保存期間と根拠法令の一覧",
                              "42": "グループ各社の勘定科目の一覧"}.get(key, "")
        regs.append(reg)
        print(f"{key}\t{reg['scope']}\t{name}\t{sum(1 for s in secs if s['reg'] == rid)}条")

    # Excel(スプレッドシート取り込み用)
    wb = openpyxl.Workbook()
    def sheet(title, header, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = w
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.number_format = "@"
                c.alignment = Alignment(vertical="top", wrap_text=False)
        return ws
    wb.remove(wb.active)
    sheet("規定", ["id", "番号", "規定名", "分類", "所管", "公開範囲", "要約", "ファイル名", "種類"],
          [[r["id"], r["no"], r["name"], r["category"], r["dept"], r["scope"], r["summary"], r["file"], r["kind"]] for r in regs],
          [6, 6, 36, 22, 14, 8, 60, 40, 6])
    sheet("条文", ["規定id", "順番", "章", "条", "見出し", "本文"],
          [[s["reg"], str(s["seq"]), s["chapter"], s["article"], s["title"], s["body"]] for s in secs],
          [6, 6, 24, 10, 24, 100])
    sheet("表", ["規定id", "シート", "行", "セル(JSON)"],
          [[t["reg"], t["sheet"], str(t["row"]), json.dumps(t["cells"], ensure_ascii=False)] for t in tables],
          [6, 20, 6, 100])
    sheet("言い換え", ["検索語", "言い換え(カンマ区切り)"], [list(x) for x in SYNONYMS], [16, 60])
    sheet("よくある質問", ["ボタン名", "検索語", "公開範囲"], [list(x) for x in FAQ], [24, 24, 10])
    built = datetime.datetime.now().strftime("%Y/%m/%d")
    sheet("情報", ["項目", "値"], [["データ作成日", built], ["元フォルダ", str(src)],
                                   ["規定数", str(len(regs))], ["条文数", str(len(secs))]], [16, 90])
    xlsx = out_dir / "規定データ.xlsx"
    wb.save(str(xlsx))

    data = {
        "builtAt": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "regs": regs, "sections": secs, "tables": tables,
        "synonyms": [{"word": a, "alts": b} for a, b in SYNONYMS],
        "faq": [{"label": a, "query": b, "scope": c} for a, b, c in FAQ],
    }
    (out_dir / "data.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    print(f"規定 {len(regs)} / 条文 {len(secs)} / 表の行 {len(tables)} → {xlsx}")


if __name__ == "__main__":
    main()
